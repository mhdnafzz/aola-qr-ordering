import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

os.makedirs("database_templates", exist_ok=True)

header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="F9FAFB")
align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")

def style_sheet(ws, headers, rows):
    ws.views.sheetView[0].showGridLines = True
    ws.append(headers)
    
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
    
    for row in rows:
        ws.append(row)
        
    for row_idx in range(2, len(rows) + 2):
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = align_left

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

# 1. Generate AOLA_Master_DB.xlsx
wb_master = openpyxl.Workbook()
ws_master = wb_master.active
ws_master.title = "Shops"

master_headers = ["Shop_ID", "Shop_Name", "Logo_URL", "Sheet_ID", "Owner_PIN", "Status"]
master_rows = [
    [
        "bismillah",
        "Bismillah Biryani & Kebabs",
        "https://images.unsplash.com/photo-1589302168068-964664d93dc0?w=200",
        "PASTE_CLIENT_SHEET_ID_HERE",
        "1234",
        "Active"
    ],
    [
        "kfc",
        "Kentucky Crisp Chicken",
        "https://images.unsplash.com/photo-1513639776629-7b61b0ac49cb?w=200",
        "PASTE_KFC_SHEET_ID_HERE",
        "4321",
        "Active"
    ]
]
style_sheet(ws_master, master_headers, master_rows)
wb_master.save("database_templates/AOLA_Master_DB.xlsx")
print("Created database_templates/AOLA_Master_DB.xlsx")

# 2. Generate Client_Template_Orders_and_Menu.xlsx
wb_client = openpyxl.Workbook()

# Tab 1: Menu
ws_menu = wb_client.active
ws_menu.title = "Menu"
menu_headers = ["Dish Name", "Price", "In Stock", "Category"]
menu_rows = [
    ["Special Chicken Dum Biryani", 240, True, "Biryani"],
    ["Mutton Hyderabadi Biryani", 320, True, "Biryani"],
    ["Chicken Tikka Kebab (6 pcs)", 190, True, "Starters"],
    ["Mutton Seekh Kebab", 230, True, "Starters"],
    ["Paneer Butter Masala", 180, True, "Curries"],
    ["Butter Garlic Naan", 45, True, "Breads"],
    ["Gulab Jamun (2 pcs)", 60, True, "Desserts"],
    ["Royal Mango Lassi", 75, False, "Beverages"]
]
style_sheet(ws_menu, menu_headers, menu_rows)

# Tab 2: Orders
ws_orders = wb_client.create_sheet(title="Orders")
orders_headers = ["Order_ID", "Timestamp", "Table_Number", "Items", "Total_Price", "Status"]
orders_rows = [
    ["ORD-1001", "2026-10-20 07:30:00", 4, "1x Special Chicken Dum Biryani (₹240)", 240, "Completed"]
]
style_sheet(ws_orders, orders_headers, orders_rows)

# Tab 3: Expenses
ws_expenses = wb_client.create_sheet(title="Expenses")
expenses_headers = ["Expense_ID", "Timestamp", "Category", "Description", "Amount"]
expenses_rows = [
    ["EXP-20261020-001", "2026-10-20 09:30:00", "Inventory", "Fresh Vegetables & Herbs", 1250],
    ["EXP-20261020-002", "2026-10-20 11:00:00", "Staff", "Kitchen Helper Daily Wages", 600],
    ["EXP-20261020-003", "2026-10-20 16:15:00", "Utilities", "Commercial LPG Gas Cylinder", 1850]
]
style_sheet(ws_expenses, expenses_headers, expenses_rows)

wb_client.save("database_templates/Client_Template_Orders_and_Menu.xlsx")
print("Created database_templates/Client_Template_Orders_and_Menu.xlsx")
